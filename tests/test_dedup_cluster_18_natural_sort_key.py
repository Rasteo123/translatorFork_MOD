# -*- coding: utf-8 -*-
"""Тесты dedup cluster-18: natural_sort_key (ranobelib/utils.py <-> gemini_reader_v3.py).

(а) Характеризационные тесты канонической реализации
    gemini_translator.utils.text_sort.natural_sort_key.
(б) Тест-маршрутизация: gemini_reader_v3._reader_parse_zip_docx_chapters
    обязан сортировать файлы внутри ZIP через каноническую функцию, а не
    через собственную копию.
"""

import io
import os
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from gemini_translator.utils.text_sort import natural_sort_key


# --- (а) характеризационные тесты канонической реализации ---

def test_natural_sort_orders_numeric_suffix_numerically():
    items = ["Ch10", "Ch2", "Ch1"]
    assert sorted(items, key=natural_sort_key) == ["Ch1", "Ch2", "Ch10"]


def test_natural_sort_is_case_insensitive_on_text_tokens():
    items = ["Bfile", "afile"]
    assert sorted(items, key=natural_sort_key) == ["afile", "Bfile"]


def test_natural_sort_handles_multiple_numeric_groups():
    items = ["v2ch10", "v2ch2", "v10ch1"]
    assert sorted(items, key=natural_sort_key) == ["v2ch2", "v2ch10", "v10ch1"]


def test_natural_sort_handles_no_digits():
    items = ["banana", "Apple", "cherry"]
    assert sorted(items, key=natural_sort_key) == ["Apple", "banana", "cherry"]


def test_natural_sort_handles_empty_string():
    assert natural_sort_key("") == [""]


def test_natural_sort_key_matches_expected_token_structure():
    assert natural_sort_key("Chapter10.docx") == ["chapter", 10, ".docx"]


# --- (б) тест-маршрутизация: вызывающий код обязан идти через каноническую функцию ---

@pytest.fixture
def reader_module(monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    import gemini_reader_v3 as reader
    return reader


def test_docx_zip_import_routes_sorting_through_canonical_natural_sort_key(reader_module, monkeypatch, tmp_path):
    """До рефакторинга это место вызова использовало собственную копию
    (_reader_natural_sort_key), а не gemini_translator.utils.text_sort.natural_sort_key.
    Подменяем каноническую функцию и убеждаемся, что подмена реально влияет
    на порядок разобранных docx-файлов внутри ZIP.
    """
    if reader_module.Document is None:
        pytest.skip("python-docx недоступен в тестовом окружении")

    from docx import Document as RealDocument

    names = ["b2.docx", "a10.docx", "a2.docx"]
    zip_path = tmp_path / "chapters.zip"

    with zipfile.ZipFile(zip_path, "w") as zf:
        for name in names:
            doc_buf = io.BytesIO()
            d = RealDocument()
            d.add_paragraph(f"text-{name}")
            d.save(doc_buf)
            zf.writestr(name, doc_buf.getvalue())

    calls = []

    def fake_natural_sort_key(value):
        calls.append(value)
        # Инвертируем обычный порядок, чтобы разница была наблюдаема.
        return [-ord(c) for c in value]

    monkeypatch.setattr(
        "gemini_translator.utils.text_sort.natural_sort_key",
        fake_natural_sort_key,
    )

    chapters = reader_module._reader_parse_zip_docx_chapters(str(zip_path))

    assert calls, (
        "_reader_parse_zip_docx_chapters должен сортировать через "
        "gemini_translator.utils.text_sort.natural_sort_key, а не через "
        "собственную копию алгоритма"
    )
    assert set(calls) == set(names)
    assert len(chapters) == len(names)
    # Подменённый ключ сортирует по убыванию кодов символов имени файла, что даёт
    # порядок b2.docx, a2.docx, a10.docx (см. проверку выше). Заголовки берутся
    # из имени файла без расширения (см. _reader_fallback_title): b2, a2, a10.
    # Если бы подмена не влияла на порядок (no-op ключ), результат совпал бы
    # с исходным порядком по namelist() ZIP: b2, a10, a2 — тест отличил бы это.
    assert [chapter.title for chapter in chapters] == ["b2", "a2", "a10"]


def test_docx_zip_import_raises_clear_error_when_text_sort_module_unavailable(reader_module, monkeypatch, tmp_path):
    """Регресс: gemini_translator.utils.text_sort — опциональный импорт (как и все прочие
    импорты из gemini_translator в этом файле), и при его недоступности вызов не должен
    падать 'AttributeError: NoneType object has no attribute natural_sort_key' —
    ожидается штатный RuntimeError, как и при отсутствии python-docx.
    """
    if reader_module.Document is None:
        pytest.skip("python-docx недоступен в тестовом окружении")

    from docx import Document as RealDocument

    zip_path = tmp_path / "chapters.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        doc_buf = io.BytesIO()
        d = RealDocument()
        d.add_paragraph("text")
        d.save(doc_buf)
        zf.writestr("a1.docx", doc_buf.getvalue())

    monkeypatch.setattr(reader_module, "_text_sort", None)

    with pytest.raises(RuntimeError):
        reader_module._reader_parse_zip_docx_chapters(str(zip_path))
